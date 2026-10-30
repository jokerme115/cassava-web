#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成小程序端软件著作权登记申请表"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def set_cell_text(cell, text, bold=False, font_size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.size = Pt(font_size)
    run.bold = bold
    run.font.name = '宋体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def set_cell_background(cell, color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def create_form():
    doc = Document()
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('计算机软件著作权登记申请表')
    run.font.size = Pt(16)
    run.bold = True
    run.font.name = '黑体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')

    doc.add_paragraph()

    # 基本信息表
    table1 = doc.add_table(rows=14, cols=2)
    table1.style = 'Table Grid'
    for row in table1.rows:
        row.cells[0].width = Cm(3)
        row.cells[1].width = Cm(13)

    set_cell_background(table1.rows[0].cells[0], 'D9E1F2')
    set_cell_text(table1.rows[0].cells[0], '软件基本信息', bold=True, font_size=10)
    table1.rows[0].cells[1].merge(table1.rows[0].cells[1])

    set_cell_text(table1.rows[1].cells[0], '软件名称', bold=True)
    set_cell_text(table1.rows[1].cells[1], '木薯腐烂面积精准量化平台（微信小程序端）')

    set_cell_text(table1.rows[2].cells[0], '版本号', bold=True)
    set_cell_text(table1.rows[2].cells[1], 'V1.0')

    set_cell_text(table1.rows[3].cells[0], '软件简称', bold=True)
    set_cell_text(table1.rows[3].cells[1], '木薯PPD小程序')

    set_cell_text(table1.rows[4].cells[0], '分类号', bold=True)
    set_cell_text(table1.rows[4].cells[1], '')

    set_cell_text(table1.rows[5].cells[0], '软件作品说明', bold=True)
    set_cell_text(table1.rows[5].cells[1], '☑ 原创    ○ 修改(含翻译软件、合成软件)')

    set_cell_text(table1.rows[6].cells[0], '软件分类', bold=True)
    set_cell_text(table1.rows[6].cells[1], '☑ 应用软件  ○ 嵌入式软件  ○ 中间件  ○ 操作系统')

    set_cell_text(table1.rows[7].cells[0], '开发完成日期', bold=True)
    set_cell_text(table1.rows[7].cells[1], '2026-01-12')

    set_cell_text(table1.rows[8].cells[0], '发表状态', bold=True)
    set_cell_text(table1.rows[8].cells[1], '○ 已发表\n☑ 未发表（软件未公之于众）')

    set_cell_text(table1.rows[9].cells[0], '首次发表日期', bold=True)
    set_cell_text(table1.rows[9].cells[1], '')

    set_cell_text(table1.rows[10].cells[0], '首次发表地点', bold=True)
    set_cell_text(table1.rows[10].cells[1], '')

    set_cell_text(table1.rows[11].cells[0], '开发方式', bold=True)
    set_cell_text(table1.rows[11].cells[1], '☑ 独立开发  ○ 合作开发  ○ 委托开发  ○ 下达任务开发')

    set_cell_text(table1.rows[12].cells[0], '著作权人', bold=True)
    set_cell_text(table1.rows[12].cells[1], '（由用户填写）')

    set_cell_text(table1.rows[13].cells[0], '权利说明', bold=True)
    set_cell_text(table1.rows[13].cells[1], '☑ 原始取得  ○ 继受取得\n☑ 全部权利  ○ 部分权利')

    doc.add_paragraph()

    # 功能和技术特点表
    table2 = doc.add_table(rows=14, cols=2)
    table2.style = 'Table Grid'
    for row in table2.rows:
        row.cells[0].width = Cm(3.5)
        row.cells[1].width = Cm(12.5)

    set_cell_background(table2.rows[0].cells[0], 'D9E1F2')
    set_cell_text(table2.rows[0].cells[0], '软件功能和\n技术特点', bold=True, font_size=9)
    table2.rows[0].cells[1].merge(table2.rows[0].cells[1])

    set_cell_text(table2.rows[1].cells[0], '开发的硬件环境', bold=True)
    set_cell_text(table2.rows[1].cells[1], '电脑（CPU: Intel/AMD多核处理器；内存: 16GB及以上）')

    set_cell_text(table2.rows[2].cells[0], '运行的硬件环境', bold=True)
    set_cell_text(table2.rows[2].cells[1], '移动端：智能手机（微信小程序运行环境）；后端：Windows/Linux服务器')

    set_cell_text(table2.rows[3].cells[0], '开发该软件的操作系统', bold=True)
    set_cell_text(table2.rows[3].cells[1], 'Windows 10/11')

    set_cell_text(table2.rows[4].cells[0], '软件开发环境或开发工具', bold=True)
    set_cell_text(table2.rows[4].cells[1], '微信开发者工具、Taro、React、TypeScript、Vite、Django')

    set_cell_text(table2.rows[5].cells[0], '软件的运行平台或操作系统', bold=True)
    set_cell_text(table2.rows[5].cells[1], '微信小程序（iOS/Android）、后端Windows Server/Linux Server')

    set_cell_text(table2.rows[6].cells[0], '软件运行支撑环境或支持软件', bold=True)
    set_cell_text(table2.rows[6].cells[1], '微信App、Python 3.10+、Django 5.2、MySQL 8.0、PyTorch深度学习框架')

    set_cell_text(table2.rows[7].cells[0], '编程语言', bold=True)
    set_cell_text(table2.rows[7].cells[1], 'TypeScript/JavaScript（前端）、Python（后端）')

    set_cell_text(table2.rows[8].cells[0], '源程序量', bold=True)
    set_cell_text(table2.rows[8].cells[1], '约12742行（前端TS/TSX代码），后端约18646行Python代码')

    set_cell_text(table2.rows[9].cells[0], '开发目的', bold=True)
    set_cell_text(table2.rows[9].cells[1], '面向木薯采后腐烂研究与实验展示，通过微信小程序端实现木薯横截面图像的快速采集与智能分割，输出PPD（采后生理衰退）量化指标，为田间现场检测提供便捷工具。')

    set_cell_text(table2.rows[10].cells[0], '面向领域/行业', bold=True)
    set_cell_text(table2.rows[10].cells[1], '农业智能检测、木薯研究、农产品品质评估、微信小程序应用')

    set_cell_text(table2.rows[11].cells[0], '软件的主要功能', bold=True)
    set_cell_text(table2.rows[11].cells[1], '1. 用户认证：支持账号注册、登录、权限管理\n2. 图像采集：支持单张拍照、批量拍照和相册补充，最多一次提交9张图片\n3. 智能分割：调用后端DeepLabV3+深度学习模型自动识别并分割木薯横截面腐烂区域\n4. PPD预测：计算腐烂面积与整图面积的比值，输出PPD量化指标\n5. 结果查看：支持原图、分割结果图、PPD数值的同步展示与复核\n6. 历史记录：本地记录预测历史，方便追溯与对比分析\n7. 参数设置：支持调整置信度阈值、选择不同模型、切换处理档案')

    set_cell_text(table2.rows[12].cells[0], '软件的技术特点', bold=True)
    set_cell_text(table2.rows[12].cells[1], '☑ 教育软件  ☑ 云计算软件\n本小程序采用Taro跨端框架结合React构建，前端实现多端代码统一。后端采用Django REST框架提供API服务。核心分割模型采用DeepLabV3+架构搭配ResNet101骨干网络，基于PyTorch深度学习框架实现。平台支持滑动窗口推理策略应对任意尺寸输入，提供移动端便捷的图像采集与实时预测能力，适用于农业现场检测场景。')

    set_cell_text(table2.rows[13].cells[0], '备注', bold=True)
    set_cell_text(table2.rows[13].cells[1], '本软件为独立开发，研究团队：海南师范大学龙海侠团队实验室；联合平台：国家木薯种质资源圃（儋州）木薯研究中心。')

    doc.add_paragraph()

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = note.add_run('※ 著作权人信息（姓名/名称、证件类型、证件号码、国籍、省份/城市等）请根据实际情况自行填写。')
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(128, 128, 128)

    output_path = r'C:\Users\Administrator\Desktop\work\02_农业研究\木薯图像分割\木薯网页源代码\木薯腐烂面积精准量化平台_小程序端_软件著作权登记申请表.docx'
    doc.save(output_path)
    print(f'小程序端文档已保存至: {output_path}')

if __name__ == '__main__':
    create_form()